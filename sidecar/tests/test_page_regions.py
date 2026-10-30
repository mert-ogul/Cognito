"""PDF page regions, OCR fallback, Word parse, and page render cache."""

from __future__ import annotations

import base64
import json
from pathlib import Path
import fitz
import pytest

from ingestion.chunking import chunk_semantic, merge_regions
from ingestion.page_render import render_page, render_region
from ingestion.parsing import parse_file
from ingestion.rpc_ingest import rpc_chunk_documents


def _two_column_pdf(path: Path) -> None:
    doc = fitz.open()
    page = doc.new_page(width=400, height=300)
    # Left column
    page.insert_textbox(fitz.Rect(20, 40, 180, 260), "LEFT COLUMN text about alpha. " * 8)
    # Right column
    page.insert_textbox(fitz.Rect(220, 40, 380, 260), "RIGHT COLUMN text about beta. " * 8)
    doc.save(path)
    doc.close()


def test_two_column_pdf_regions_on_chunks(tmp_path: Path) -> None:
    pdf = tmp_path / "cols.pdf"
    _two_column_pdf(pdf)
    parsed = parse_file(pdf)
    assert parsed["fileType"] == "pdf"
    assert parsed["pages"]
    page0 = parsed["pages"][0]
    assert page0["scanned"] is False
    assert len(page0["blocks"]) >= 2
    for b in page0["blocks"]:
        assert "bbox" in b and len(b["bbox"]) == 4

    drafts = chunk_semantic(parsed["text"], "pdf", pages=parsed["pages"])
    assert drafts
    with_regions = [c for c in drafts if c.regions]
    assert with_regions
    for c in with_regions:
        assert c.regions is not None
        for r in c.regions:
            assert r["page"] == 1
            assert r["x1"] > r["x0"] and r["y1"] > r["y0"]

    rpc = rpc_chunk_documents({"items": [parsed]})
    chunk0 = rpc["items"][0]["chunks"][0]
    assert chunk0["regions"] is not None
    regions = json.loads(chunk0["regions"])
    assert isinstance(regions, list) and regions


def test_merge_regions_per_page() -> None:
    raw = [
        {"page": 1, "x0": 10, "y0": 10, "x1": 50, "y1": 40},
        {"page": 1, "x0": 48, "y0": 12, "x1": 90, "y1": 42},  # near / overlapping
        {"page": 1, "x0": 10, "y0": 200, "x1": 40, "y1": 230},  # far
        {"page": 2, "x0": 0, "y0": 0, "x1": 10, "y1": 10},
    ]
    merged = merge_regions(raw, gap=4.0)
    page1 = [r for r in merged if r["page"] == 1]
    page2 = [r for r in merged if r["page"] == 2]
    assert len(page1) == 2
    assert len(page2) == 1
    top = min(page1, key=lambda r: r["y0"])
    assert top["x0"] <= 10 and top["x1"] >= 90


def test_render_region_size_and_cache_eviction(tmp_path: Path) -> None:
    pdf = tmp_path / "r.pdf"
    doc = fitz.open()
    page = doc.new_page(width=200, height=200)
    page.insert_text((20, 40), "Hello render")
    doc.save(pdf)
    doc.close()

    cache = tmp_path / "cache"
    first = render_region(pdf, 1, [10, 10, 100, 80], scale=2, cache_dir=cache)
    assert first["width"] > 0 and first["height"] > 0
    raw = base64.b64decode(first["pngBase64"])
    assert raw[:8] == b"\x89PNG\r\n\x1a\n"
    key_path = cache / f"{first['cacheKey']}.png"
    assert key_path.is_file()

    second = render_region(pdf, 1, [10, 10, 100, 80], scale=2, cache_dir=cache)
    assert second["cacheKey"] == first["cacheKey"]

    # Force tiny cache and fill with distinct clips so eviction runs.
    for i in range(8):
        render_region(
            pdf,
            1,
            [10 + i, 10, 90 + i, 80],
            scale=2,
            cache_dir=cache,
            max_cache_bytes=8_000,
        )
    remaining = list(cache.glob("*.png"))
    assert remaining
    assert len(remaining) < 8
    total = sum(p.stat().st_size for p in remaining)
    assert total <= 8_000 + max(p.stat().st_size for p in remaining)


def test_render_page_caps_width(tmp_path: Path) -> None:
    pdf = tmp_path / "wide.pdf"
    doc = fitz.open()
    doc.new_page(width=1000, height=200)
    doc.save(pdf)
    doc.close()
    out = render_page(pdf, 1, scale=2, cache_dir=tmp_path / "c")
    assert out["width"] <= 1400


def _image_only_pdf(path: Path) -> None:
    doc = fitz.open()
    page = doc.new_page(width=200, height=200)
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 40, 40), 0)
    pix.set_rect(pix.irect, (200, 200, 200))
    page.insert_image(fitz.Rect(20, 20, 180, 180), pixmap=pix)
    doc.save(path)
    doc.close()


def test_scanned_page_ocr_mocked(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    pdf = tmp_path / "scan.pdf"
    _image_only_pdf(pdf)

    class FakeOCR:
        def __init__(self, *_a, **_k):
            pass

        def recognize(self):
            # normalized [x, y, w, h] bottom-left origin
            return [("Scanned line one", 0.9, [0.1, 0.7, 0.5, 0.1])]

    import sys
    import types

    fake_mod = types.ModuleType("ocrmac.ocrmac")
    fake_mod.OCR = FakeOCR  # type: ignore[attr-defined]
    pkg = types.ModuleType("ocrmac")
    pkg.ocrmac = fake_mod  # type: ignore[attr-defined]
    monkeypatch.setitem(sys.modules, "ocrmac", pkg)
    monkeypatch.setitem(sys.modules, "ocrmac.ocrmac", fake_mod)
    monkeypatch.setattr("ingestion.parsing.platform.system", lambda: "Darwin")

    parsed = parse_file(pdf)
    assert parsed["scanStatus"] == "ocr"
    assert parsed.get("textSource") == "ocr"
    assert "Scanned line one" in parsed["text"]
    assert parsed["pages"][0]["ocr"] is True
    assert parsed["pages"][0]["blocks"]


def test_scanned_page_no_text_non_mac(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    pdf = tmp_path / "scan2.pdf"
    _image_only_pdf(pdf)

    monkeypatch.setattr("ingestion.parsing.platform.system", lambda: "Linux")
    parsed = parse_file(pdf)
    assert parsed["scanStatus"] == "no_text"
    assert parsed["pages"][0]["scanned"] is True
    assert parsed["pages"][0]["ocr"] is False
    assert parsed["pages"][0]["text"] == ""


def test_parse_docx_headings_and_table(tmp_path: Path) -> None:
    from docx import Document

    path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_heading("Intro Title", level=1)
    doc.add_heading("Details", level=2)
    doc.add_paragraph("Body paragraph here.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Role"
    table.cell(1, 0).text = "Ada"
    table.cell(1, 1).text = "Author"
    doc.save(path)

    parsed = parse_file(path)
    assert parsed["fileType"] == "docx"
    assert parsed["pageCount"] is None
    assert parsed["scanStatus"] is None
    assert "# Intro Title" in parsed["text"] or "Intro Title" in parsed["text"]
    assert "Details" in parsed["text"]
    assert "Body paragraph here." in parsed["text"]
    assert "Name\tRole" in parsed["text"] or "Ada\tAuthor" in parsed["text"]

    chunks = chunk_semantic(parsed["text"], "docx")
    assert chunks
    assert any(c.regions is None for c in chunks)
    assert any("Intro Title" in c.text or "Details" in " > ".join(c.heading_trail) for c in chunks)
